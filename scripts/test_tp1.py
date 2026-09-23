# -*- coding: utf-8 -*-
"""
scripts/test_tp1.py — Suite de Pruebas Automatizadas y Verificación de Integridad
Cátedra: Gestión y Mantenimiento Electromecánico — UTN FRSR (Año 2026)

Valida con rigor de ingeniería y QA:
1. Compilación estricta py_compile de app.py y scripts/report_builder.py.
2. Generación dinámica de informe técnico Word (.docx) con soporte multigrupal y Standby motor.
3. Existencia e integridad de los entregables consolidados en la raíz de TP N°1/.
4. Existencia de los enunciados oficiales en documentos_catedra/.
5. Limpieza de redundancias en la carpeta padre Gestión de mantenimiento/.
6. Ejecución headless completa de la aplicación Streamlit (AppTest) con 0 excepciones.
"""

import os
import sys
import pathlib
import unittest
import py_compile
import docx

# Directorios de referencia
TP1_ROOT = pathlib.Path(__file__).resolve().parent.parent
PARENT_DIR = TP1_ROOT.parent
SCRIPTS_DIR = TP1_ROOT / "scripts"

# Asegurar importación de scripts/
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

import report_builder


class TestTP1EngineeringPlatform(unittest.TestCase):
    """Batería de pruebas para la plataforma de mantenimiento del TP N°1."""

    def test_01_py_compile_syntax(self):
        """Verifica que app.py y report_builder.py compilen sin errores de sintaxis."""
        app_path = TP1_ROOT / "app.py"
        builder_path = SCRIPTS_DIR / "report_builder.py"

        self.assertTrue(app_path.exists(), "app.py debe existir en la raíz de TP N°1")
        self.assertTrue(builder_path.exists(), "report_builder.py debe existir en scripts/")

        # Compilar app.py
        compiled_app = py_compile.compile(str(app_path), doraise=True)
        self.assertTrue(os.path.exists(compiled_app), "Fallo la compilación de app.py")

        # Compilar scripts/report_builder.py
        compiled_builder = py_compile.compile(str(builder_path), doraise=True)
        self.assertTrue(os.path.exists(compiled_builder), "Fallo la compilación de report_builder.py")
        print("\n[OK] Test 1: Compilación py_compile completada exitosamente.")

    def test_02_report_builder_multi_student_and_standby(self):
        """Valida que report_builder genere un documento Word válido > 20 KB con múltiples alumnos y standby activo."""
        students_test = [
            {"nombre": "Martín", "apellido": "Méndez", "legajo": "11335"},
            {"nombre": "Juan", "apellido": "Pérez", "legajo": "11400"},
            {"nombre": "Gonzalo", "apellido": "García", "legajo": "11450"}
        ]

        # 1. Escenario con Standby inactivo
        buf_std_off = report_builder.build_technical_report(students=students_test, standby_active=False)
        size_off = len(buf_std_off.getvalue())
        self.assertGreater(size_off, 20480, f"El documento Word (standby off) debe superar los 20 KB (obtenido: {size_off} bytes)")

        # 2. Escenario con Standby activo
        buf_std_on = report_builder.build_technical_report(students=students_test, standby_active=True)
        size_on = len(buf_std_on.getvalue())
        self.assertGreater(size_on, 20480, f"El documento Word (standby on) debe superar los 20 KB (obtenido: {size_on} bytes)")

        # 3. Validar contenido y tablas en el docx
        doc = docx.Document(buf_std_on)
        self.assertGreater(len(doc.paragraphs), 30, "El documento debe contener más de 30 párrafos de texto técnico")
        self.assertGreaterEqual(len(doc.tables), 8, "El documento debe contener al menos 8 tablas estructuradas")

        # Verificar presencia de los alumnos en la tabla de portada
        tbl_roster = doc.tables[0]
        roster_text = " ".join([cell.text for row in tbl_roster.rows for cell in row.cells])
        self.assertIn("Méndez Martín", roster_text, "Martín Méndez debe figurar en la portada")
        self.assertIn("11335", roster_text, "Legajo 11335 debe figurar en la portada")
        self.assertIn("Pérez Juan", roster_text, "Juan Pérez debe figurar en la portada")
        self.assertIn("11400", roster_text, "Legajo 11400 debe figurar en la portada")

        # Verificar presencia de secciones clave (Ishikawa 6M, 5 Whys, Standby)
        full_doc_text = " ".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
        self.assertIn("Ishikawa", full_doc_text, "Debe contener sección de Ishikawa")
        self.assertIn("Maquinaria", full_doc_text, "Debe contener dimensión Maquinaria")
        self.assertIn("5 Porqués", full_doc_text, "Debe contener metodología de 5 Porqués")
        self.assertIn("Standby", full_doc_text, "Debe contener análisis de motor Standby")
        self.assertIn("95,000", full_doc_text, "Debe reflejar el ahorro de $95,000 USD")
        self.assertIn("16.40", full_doc_text, "Debe reflejar el MTTR global de 16.40 h")
        print(f"[OK] Test 2: Generación Word con multigrupo y Standby verificada ({size_on / 1024:.1f} KB).")

    def test_03_root_deliverables_exist(self):
        """Verifica que los entregables requeridos existan en la raíz de TP N°1."""
        docx_file = TP1_ROOT / "TP1_Informe_Tecnico_Electromecanico_UTN.docx"
        xlsx_file = TP1_ROOT / "TP1_Gestion_Mantenimiento_UTN.xlsx"
        ipynb_file = TP1_ROOT / "TP1_Gestion_Mantenimiento_UTN.ipynb"
        app_file = TP1_ROOT / "app.py"

        self.assertTrue(docx_file.exists(), "Falta TP1_Informe_Tecnico_Electromecanico_UTN.docx en raíz")
        self.assertGreater(docx_file.stat().st_size, 20480, "DOCX debe pesar más de 20 KB")

        self.assertTrue(xlsx_file.exists(), "Falta TP1_Gestion_Mantenimiento_UTN.xlsx en raíz")
        self.assertGreater(xlsx_file.stat().st_size, 10240, "XLSX debe pesar más de 10 KB")

        self.assertTrue(ipynb_file.exists(), "Falta TP1_Gestion_Mantenimiento_UTN.ipynb en raíz")
        self.assertGreater(ipynb_file.stat().st_size, 10240, "IPYNB debe pesar más de 10 KB")

        self.assertTrue(app_file.exists(), "Falta app.py en raíz")
        print("[OK] Test 3: Entregables consolidados en la raíz de TP N°1 verificados.")

    def test_04_documentos_catedra_directory(self):
        """Verifica que documentos_catedra/ contenga los enunciados oficiales."""
        cat_dir = TP1_ROOT / "documentos_catedra"
        self.assertTrue(cat_dir.exists() and cat_dir.is_dir(), "documentos_catedra/ debe existir")

        docx_cat = cat_dir / "TP1_GESTIÓN Y MANTENIMIENTO ELECTROMECÁNICO.docx"
        pdf_cat = cat_dir / "TP1_GESTIÓN Y MANTENIMIENTO ELECTROMECÁNICO.pdf"

        self.assertTrue(docx_cat.exists(), "Falta enunciado .docx en documentos_catedra/")
        self.assertGreater(docx_cat.stat().st_size, 500000, "Enunciado .docx debe superar 500 KB")

        self.assertTrue(pdf_cat.exists(), "Falta enunciado .pdf en documentos_catedra/")
        self.assertGreater(pdf_cat.stat().st_size, 100000, "Enunciado .pdf debe superar 100 KB")
        print("[OK] Test 4: Enunciados de cátedra en documentos_catedra/ verificados.")

    def test_05_parent_directory_cleanup(self):
        """Verifica la limpieza de archivos redundantes en la carpeta padre de la asignatura."""
        redundant_files = [
            "index.html",
            "TP1_Gestion_Mantenimiento.html",
            "ORIGINAL_REQUEST.md",
            "TP1_GESTIÓN Y MANTENIMIENTO ELECTROMECÁNICO_RESUELTO.docx"
        ]
        for fname in redundant_files:
            fpath = PARENT_DIR / fname
            self.assertFalse(fpath.exists(), f"El archivo {fname} no debe existir en la carpeta padre")

        # Asegurar preservación de bibliotecas de la cátedra
        herramientas_dir = PARENT_DIR / "Herramientas de análisis de problemas"
        normas_dir = PARENT_DIR / "Normas"
        self.assertTrue(herramientas_dir.exists(), "Herramientas de análisis... debe preservarse")
        self.assertTrue(normas_dir.exists(), "Normas/ debe preservarse")
        print("[OK] Test 5: Limpieza de carpeta padre y preservación de bibliotecas verificada.")

    def test_06_streamlit_headless_apptest(self):
        """Verifica la ejecución de la app Streamlit en modo headless sin excepciones."""
        try:
            from streamlit.testing.v1 import AppTest
        except ImportError:
            self.skipTest("streamlit.testing.v1.AppTest no está disponible en este entorno")

        app_file = str(TP1_ROOT / "app.py")
        at = AppTest.from_file(app_file, default_timeout=35)
        at.run()

        self.assertEqual(len(at.exception), 0, f"AppTest inicial arrojó excepciones: {at.exception}")

        # Probar toggle de motor Standby
        standby_toggle = None
        for t in at.toggle:
            if "Standby" in t.label:
                standby_toggle = t
                break
        self.assertIsNotNone(standby_toggle, "Toggle de motor Standby debe existir en la barra lateral")
        standby_toggle.set_value(True)
        at.run()
        self.assertEqual(len(at.exception), 0, f"AppTest con Standby activo arrojó excepciones: {at.exception}")

        # Probar toggle de repeticiones N_est
        rep_toggle = None
        for t in at.toggle:
            if "repetición" in t.label:
                rep_toggle = t
                break
        self.assertIsNotNone(rep_toggle, "Toggle de repetición N_est debe existir")
        rep_toggle.set_value(True)
        at.run()
        self.assertEqual(len(at.exception), 0, f"AppTest con repeticiones N_est arrojó excepciones: {at.exception}")

        print("[OK] Test 6: Streamlit headless AppTest ejecutado con 0 excepciones en todos los estados.")


if __name__ == "__main__":
    print("=" * 75)
    print(" EJECUTANDO SUITE DE PRUEBAS AUTOMATIZADAS TP N°1 (UTN FRSR)")
    print("=" * 75)
    unittest.main(verbosity=2)
